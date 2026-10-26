"""Word（.docx）loader：python-docx 解出标题层级、正文、表格与图片。

**为什么不走 OCR**：docx 的结构是明确写在文件里的 —— 段落的样式就叫 ``Heading 1``，
不需要识别，也不该花 GPU。这条路的出口和 OCR 一样是 **markdown**：
下游 `markdown-semantic` 切分器只认 markdown 标题，两条路必须在出口对齐。

**图片**用和 PDF 那条路同一套机制（`ingest/assets.py`）：本体进资产目录、正文里留
`![image](asset:…)` 原位标记。这样它自动接上后面那套「图 + 上下文 → VLM 描述」的流程。
实测一份 10MB 的文档里有 45 张内容图 —— 不取的话它们会**静默消失**。
"""

from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterator
from pathlib import Path

from docx import Document as open_docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from rag.config.schema import IngestConfig
from rag.contracts import Document, ImageAsset, Position
from rag.ingest.assets import asset_store_for_cfg, make_marker
from rag.ingest.loaders.base import (
    doc_id_for,
    mime_for,
    normalize_newlines,
    register_loader,
    rel_uri,
    to_markdown_table,
)

SUFFIXES = (".docx",)

# 样式名随文档语言变：英文是 "Heading 1"，中文 Word 存的是 "标题 1"。
# 两种都认 —— 只认一种的话，中文文档会被整篇当成正文，标题层级全丢。
_HEADING = re.compile(r"^(?:heading|标题)\s*([1-6])$", re.IGNORECASE)

# 直接格式推断标题时，超过这个长度的一律不算 —— 实测正文里最长的加粗段有 4307 字，
# 那是强调不是标题
HEADING_MAX_CHARS = 60


def _blocks(document: object) -> Iterator[Paragraph | Table]:
    """按**文档顺序**产出段落与表格。

    python-docx 把 ``document.paragraphs`` 和 ``document.tables`` 分成两个集合，
    直接用它们会丢掉相对顺序 —— 表格全跑到文末，语义就断了。
    """
    for child in document.element.body.iterchildren():  # type: ignore[attr-defined]
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)  # type: ignore[arg-type]
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)  # type: ignore[arg-type]


def _heading_level(paragraph: Paragraph) -> int | None:
    """样式给出的标题级别。

    **样式可能是 None**：实测一份 10MB 的文档里 890 个段落全是 None（整篇没用样式），
    直接取 ``.name`` 会抛 AttributeError，把整份文档变成 skipped、原因还是一句
    ``'NoneType' object has no attribute 'name'`` —— 照着修都没法修。
    """
    name = getattr(paragraph.style, "name", "") or ""
    match = _HEADING.match(name.strip())
    return int(match.group(1)) if match else None


_BLIP = qn("a:blip")  # 内联图与浮动图的图片引用都在这个元素上
_EMBED = qn("r:embed")


def _image_ids(paragraph: Paragraph) -> list[str]:
    """段落里的图片 rId，按出现顺序。"""
    ids = [blip.get(_EMBED) for blip in paragraph._p.findall(f".//{_BLIP}")]
    return [rid for rid in ids if rid]


def _uniform_size(runs: list) -> float | None:
    """段落里所有 run 用同一个字号时返回它，否则 None（混字号的多半是正文）。"""
    sizes = {run.font.size.pt for run in runs if run.font.size}
    return sizes.pop() if len(sizes) == 1 else None


def _body_font_size(docx: object) -> float | None:
    """正文基准字号 = 按字符数最多的那个。整篇没有字号信息时返回 None。"""
    weight: Counter[float] = Counter()
    for paragraph in docx.paragraphs:  # type: ignore[attr-defined]
        for run in paragraph.runs:
            if run.text.strip() and run.font.size:
                weight[run.font.size.pt] += len(run.text.strip())
    return weight.most_common(1)[0][0] if weight else None


def _formatted_heading_candidate(paragraph: Paragraph, body_size: float) -> float | None:
    """整段加粗 + 比正文大 + 够短 → 返回它的字号，否则 None。"""
    text = paragraph.text.strip()
    if not text or len(text) > HEADING_MAX_CHARS:
        return None
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs or not all(run.font.bold for run in runs):
        return None
    size = _uniform_size(runs)
    return size if size is not None and size > body_size else None


def _formatted_heading_levels(docx: object, body_size: float) -> dict[float, int]:
    """字号 → 层级（大的级别高）。

    **只在整篇没用样式时才走这条路。** 实测一份 10MB 的文档里 890 个段落全是
    `style=None` —— 作者用字号和加粗表达结构，一个 Heading 样式都没用。这种情况
    按样式找标题会一个都找不到，整篇文档变成没有层级的平文本。
    """
    sizes = {
        size
        for paragraph in docx.paragraphs  # type: ignore[attr-defined]
        if (size := _formatted_heading_candidate(paragraph, body_size))
    }
    return {size: rank + 1 for rank, size in enumerate(sorted(sizes, reverse=True))}


def _has_styled_headings(docx: object) -> bool:
    return any(_heading_level(p) for p in docx.paragraphs)  # type: ignore[attr-defined]


@register_loader(*SUFFIXES)
class DocxLoader:
    """一段 = 一行；标题补 ``#``；表格转竖线表。"""

    name = "docx"
    suffixes = SUFFIXES

    def __init__(
        self, cfg: IngestConfig, ocr: object | None = None, assets: object | None = None
    ) -> None:
        """签名见 base.build_loader 的统一约定：这条路上用不到 OCR，图片要落到 `assets`。"""
        self._cfg = cfg
        self._assets = assets

    def load(self, path: Path, *, root: Path | None = None) -> Document:
        try:
            docx = open_docx(str(path))
        except Exception as exc:
            raise ValueError(f"docx 打开失败: {exc}") from exc

        doc_id = doc_id_for(path, root)
        store = self._assets if self._assets is not None else asset_store_for_cfg(self._cfg.assets)
        lines: list[str] = []
        images: list[ImageAsset] = []
        positions: list[Position] = []
        offset = 0

        def push(text: str) -> None:
            """追加一个块，并记下它的起始偏移。

            docx **没有页的概念**（分页是 Word 渲染时算的，不在文件里），所以用
            **块在文档里的序号**定位 —— 段落、表格、图片各算一块，按文档顺序编号。
            """
            nonlocal offset
            if lines:
                offset += 2  # 块之间用 "\n\n" 连接
            positions.append(Position(offset=offset, page=0, line=len(lines) + 1))
            lines.append(text)
            offset += len(text)

        # 有样式就信样式，不去猜；整篇没用样式才按直接格式推断（见 _formatted_heading_levels）
        body_size = _body_font_size(docx)
        levels: dict[float, int] = {}
        if not _has_styled_headings(docx) and body_size is not None:
            levels = _formatted_heading_levels(docx, body_size)

        for block in _blocks(docx):
            if isinstance(block, Table):
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                if table := to_markdown_table(rows):
                    push(table)
                continue

            text = block.text.strip()
            if text:
                level = _heading_level(block)
                if level is None and levels and body_size is not None:
                    # 兜底走的是同一个判定（整段加粗 + 比正文大 + 够短）
                    size = _formatted_heading_candidate(block, body_size)
                    if size is not None:
                        level = levels.get(size)
                push(f"{'#' * level} {text}" if level else text)

            # 图片排在所在段落的文字之后。段内的精确顺序（图在文字前还是后）
            # 需要逐个 run 判，收益很小 —— 图片通常独占一段，标记落在段落之间就够了，
            # 上下游要的「前后段落」照样对得上。
            for rid in _image_ids(block):
                part = docx.part.related_parts.get(rid)
                if part is None:
                    raise ValueError(f"段落引用了图片 {rid}，但文档里找不到它")
                name = f"img{len(images) + 1:03d}{Path(str(part.partname)).suffix or '.png'}"
                uri = store.save(part.blob, doc_id=doc_id, name=name)
                marker = make_marker(uri)
                images.append(
                    ImageAsset(marker=marker, asset_uri=uri, source_name=str(part.partname))
                )
                push(marker)

        return Document(
            doc_id=doc_id,
            source_uri=rel_uri(path, root),
            mime=mime_for(path),
            # 段落之间空一行：markdown 的段落边界，也是递归切块的首选断点
            text=normalize_newlines("\n\n".join(lines)),
            images=images,
            positions=positions,
            metadata={"parser": self.name, "image_count": len(images)},
        )
