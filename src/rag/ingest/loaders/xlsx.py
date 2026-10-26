"""Excel（.xlsx）loader：openpyxl 读出每个工作表，**一行一条记录**。

**不用 pandas**：它现在被机器的应用控制策略挡着（`parsers.pyd` 加载失败），
而且这里只要「逐行读单元格」这点事，openpyxl 直接做，少一层依赖。

**为什么不输出 markdown 表格**：表格数据的形状和散文不一样 —— 一行本身就是最小的
完整语义单元，不该再按文本结构切（下游切分器会把两行并成一块、或把一行切两半）。
loader 在这里声明 ``chunk_strategy: table-rows``，让切分器按行切（见 chunkers/rows.py）。

**也不输出裸行**：``| 甲公司 | 512 |`` 不知道列名就检索不到，所以每条记录都把
表名与列名焊进去 —— ``客户清单：名称=甲公司，数量=512``。这一条是有意的：
检索时用户问的是「甲公司的数量是多少」，答案必须自带这些词。

代价：同一份表里每行都重复了一遍列名，比表格费 token。但表格那条路的块本来就
不该按大小压 —— 一条记录再长也是它自己。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from rag.config.schema import IngestConfig
from rag.contracts import Document, Position
from rag.ingest.loaders.base import (
    doc_id_for,
    mime_for,
    normalize_newlines,
    register_loader,
    rel_uri,
)

SUFFIXES = (".xlsx", ".xlsm")


def _cell(value: Any) -> str:
    """单元格 → 字符串。日期、浮点、None 都要收敛成可检索的文本。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel 里 512 读出来常是 512.0，留着小数点会让「512」检索不到
        return str(int(value))
    return str(value).strip()


def _column_names(header: list[str]) -> list[str]:
    """表头里空着的列按位置编号 —— 否则它的值会变成 ``=512`` 这种没法检索的东西。"""
    return [name or f"列{i + 1}" for i, name in enumerate(header)]


def _records(sheet: str, rows: list[tuple[int, list[str]]]) -> list[tuple[int, str]]:
    """一个工作表 → 每行一条自描述记录，带上它在表里的**原始行号**。

    行号要留着：它是 Excel 里唯一天然的定位（「客户清单第 5 行」），而
    Milvus 的 `line` 字段正需要它。
    """
    rows = [(number, cells) for number, cells in rows if any(cells)]
    if not rows:
        return []

    if len(rows) == 1:
        # 只有一行：无从知道列名（那一行更可能是数据而不是表头），按位置编号，
        # 总比按「表头」处理、结果一条记录都不产出强
        names, data = _column_names([""] * len(rows[0][1])), rows
    else:
        names, data = _column_names(rows[0][1]), rows[1:]

    out: list[tuple[int, str]] = []
    for number, row in data:
        padded = list(row) + [""] * (len(names) - len(row))
        pairs = [f"{name}={value}" for name, value in zip(names, padded, strict=False) if value]
        if pairs:
            out.append((number, f"{sheet}：" + "，".join(pairs)))
    return out


@register_loader(*SUFFIXES)
class XlsxLoader:
    name = "xlsx"
    suffixes = SUFFIXES

    def __init__(
        self, cfg: IngestConfig, ocr: object | None = None, assets: object | None = None
    ) -> None:
        """签名见 base.build_loader 的统一约定：这条路上用不到配置与 OCR。"""

    def load(self, path: Path, *, root: Path | None = None) -> Document:
        try:
            # read_only=True：整表读进内存在大文件上会爆；data_only=True：取公式的结果值
            workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f"xlsx 打开失败: {exc}") from exc

        rows_with_numbers: list[tuple[int, str]] = []
        sheet_count = len(workbook.sheetnames)
        for sheet in workbook.worksheets:
            rows = [
                (number, [_cell(cell) for cell in values])
                for number, values in enumerate(sheet.iter_rows(values_only=True), start=1)
            ]
            rows_with_numbers.extend(_records(sheet.title, rows))
        workbook.close()

        # 一条记录一行；位置用**表内行号**，那才是 Excel 里能对得上的定位
        records = [text for _, text in rows_with_numbers]
        positions: list[Position] = []
        offset = 0
        for number, text in rows_with_numbers:
            if positions:
                offset += 1  # 记录之间用 "\n" 连接
            positions.append(Position(offset=offset, page=0, line=number))
            offset += len(text)

        return Document(
            doc_id=doc_id_for(path, root),
            source_uri=rel_uri(path, root),
            mime=mime_for(path),
            # 一条记录一行；切分器按行切，正好一行一块
            text=normalize_newlines("\n".join(records)),
            positions=positions,
            metadata={
                "parser": self.name,
                "sheet_count": sheet_count,
                # 声明切法：表格是「一行一条记录」，不走 markdown-semantic
                "chunk_strategy": "table-rows",
            },
        )
