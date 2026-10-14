"""locate：切完块之后，把每块反查回「第几页 / 第几行」。

全部本地 —— OCR 用假客户端，不碰网络。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document as open_docx

from rag.config.schema import ChunkingConfig, IndexingConfig, Settings
from rag.contracts import Embeddings
from rag.ingest import prepare
from rag.ingest.loaders import load_source
from rag.ingest.ocr import OcrPage


class _FakeOcr:
    name = "paddleocr"  # 真实 client 都会报这个名字（工厂约定的一部分）

    def __init__(self, *pages: str) -> None:
        self._pages = list(pages)

    def parse(self, path: Path, *, file_type: int) -> list[OcrPage]:
        return [OcrPage(page_index=i, markdown=t) for i, t in enumerate(self._pages)]


class _StubEmbedder:
    @property
    def dim(self) -> int:
        return 2

    async def embed_documents(self, texts: list[str]) -> Embeddings:
        return Embeddings(dense=[[1.0, 0.0] for _ in texts])

    async def embed_query(self, text: str) -> Embeddings:
        return Embeddings(dense=[[1.0, 0.0]])


def _settings(*, size: int = 60, strategy: str = "recursive") -> Settings:
    return Settings(
        indexing=IndexingConfig(chunk=ChunkingConfig(strategy=strategy, size=size, overlap=0))
    )


def _image(tmp_path: Path) -> Path:
    target = tmp_path / "scan.png"
    target.write_bytes(b"not-a-real-png")
    return target


# ---- OCR 那条路：逐页打点 ----


def test_each_chunk_is_mapped_back_to_its_page(tmp_path: Path) -> None:
    """页是产线天然会告诉我们的，位置映射就该把它留下来。"""
    pages = [
        "第一页：" + "甲" * 40 + "。\n第二行。\n第三行。",
        "第二页：" + "乙" * 40 + "。",
        "第三页：" + "丙" * 40 + "。",
    ]

    prep = prepare(_image(tmp_path), _settings(), _FakeOcr(*pages), _StubEmbedder())

    assert [c.page for c in prep.chunks] == [1, 2, 3]
    assert [c.line for c in prep.chunks] == [1, 1, 1]


def test_line_number_is_right_within_a_page(tmp_path: Path) -> None:
    """行号是**页内**第几行，不是全文档第几行。"""
    page = "\n".join(f"第{i}行，写点够长的内容让它单独成块。" for i in range(1, 6))

    prep = prepare(_image(tmp_path), _settings(size=25), _FakeOcr(page), _StubEmbedder())

    assert [c.page for c in prep.chunks] == [1] * len(prep.chunks)
    assert [c.line for c in prep.chunks] == [1, 2, 3, 4, 5]


def test_locate_survives_the_splitter_rewriting_whitespace(tmp_path: Path) -> None:
    """`MarkdownHeaderTextSplitter` **会改写空白**（连续空行、行尾空格、CRLF 都归一化），
    所以切出来的块正文不是原文的逐字子串 —— 直接 find 会大面积落空（实测过）。

    所以定位前先**去掉所有空白**再匹配。这条用例专门钉住这个坑。
    """
    pages = [
        "# 标题一\n\n第一段。\n\n\n\n第二段，前面有连续空行。\n\n",
        "# 标题二\n\n第三段，带行尾空格。   \n第四段。\r\n",
    ]

    prep = prepare(
        _image(tmp_path),
        _settings(size=40, strategy="markdown"),
        _FakeOcr(*pages),
        _StubEmbedder(),
    )

    assert all(c.page for c in prep.chunks), "空白被归一化之后仍然要能定位"
    assert [c.page for c in prep.chunks] == sorted(c.page for c in prep.chunks)


def test_page_is_zero_when_the_document_has_no_positions(tmp_path: Path) -> None:
    """纯文本 loader 不产位置 —— 那时 page/line 保持 0，不假装知道。"""
    target = tmp_path / "a.md"
    target.write_text("# 标题\n\n正文。", encoding="utf-8")

    prep = prepare(target, _settings(), None)

    assert prep.chunks and all(c.page == 0 and c.line == 0 for c in prep.chunks)


# ---- docx：没有页，用块序号 ----


def test_docx_positions_use_block_ordinal(tmp_path: Path) -> None:
    """docx 里没有「页」（分页是 Word 渲染时算的，不在文件里），所以用块序号定位。"""
    target = tmp_path / "a.docx"
    doc = open_docx()
    for i in range(1, 6):
        doc.add_paragraph(f"第 {i} 段，写长一点让它有自己的位置。")
    doc.save(target)

    document = load_source(target, Settings().ingest).documents[0]

    assert [p.line for p in document.positions] == [1, 2, 3, 4, 5]
    assert all(p.page == 0 for p in document.positions)


# ---- xlsx：用表内行号 ----


def test_xlsx_positions_use_sheet_row_numbers(tmp_path: Path) -> None:
    """Excel 里唯一天然的定位是「第几行」，位置就该记这个。"""
    target = tmp_path / "a.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "客户清单"
    sheet.append(["名称", "数量"])  # 第 1 行 = 表头
    sheet.append(["甲公司", 512])  # 第 2 行
    sheet.append(["乙公司", 1024])  # 第 3 行
    book.save(target)

    document = load_source(target, Settings().ingest).documents[0]

    assert [p.line for p in document.positions] == [2, 3], "行号要含表头偏移，对得上 Excel"


def test_table_rows_chunks_carry_their_sheet_row(tmp_path: Path) -> None:
    target = tmp_path / "a.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "客户清单"
    sheet.append(["名称", "数量"])
    sheet.append(["甲公司", 512])
    book.save(target)

    prep = prepare(target, _settings(), None)

    assert len(prep.chunks) == 1
    assert prep.chunks[0].page == 0
    assert prep.chunks[0].line == 2, "第一条数据在表里的第 2 行"
