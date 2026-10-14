"""docx / html / xlsx 走代码解析：标题层级、表格、以及出口是 markdown 这件事。

这三条路的共同要求是**出口和 OCR 对齐**（都是 markdown）—— 否则下游
markdown-semantic 切分器的第一级对它们就是空转。所以每条都验「标题变成了 #」。
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pymupdf
import pytest
from docx import Document as open_docx
from docx.shared import Pt

from rag.config.schema import ChunkingConfig, IndexingConfig, IngestConfig, Settings
from rag.contracts import Embeddings
from rag.ingest import prepare
from rag.ingest.loaders import LOADERS, build_loader, load_source
from rag.ingest.loaders.base import to_markdown_table
from rag.ingest.loaders.docx import _heading_level


def _load(path: Path):
    return load_source(path, IngestConfig())


# ---- 共用的表格生成器 ----­


def test_markdown_table_pads_ragged_rows_and_escapes_pipes() -> None:
    table = to_markdown_table([["名称", "数量"], ["含|竖线", "512"], ["只有一格"]])

    assert table.splitlines()[0] == "| 名称 | 数量 |"
    assert "含\\|竖线" in table, "裸竖线会把整张表切错列"
    assert table.splitlines()[-1].endswith("|  |"), "短行要补空格，否则列数对不齐"


def test_markdown_table_skips_all_empty_rows() -> None:
    assert to_markdown_table([["", ""], ["", ""]]) == ""


# ---- docx ----­


def _docx(tmp_path: Path) -> Path:
    target = tmp_path / "report.docx"
    doc = open_docx()
    doc.add_heading("年报", level=1)
    doc.add_paragraph("总览。")
    doc.add_heading("财务", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "营收"
    table.cell(1, 1).text = "100"
    doc.add_paragraph("表后的一段。")
    doc.save(target)
    return target


def test_docx_headings_become_markdown_headings(tmp_path: Path) -> None:
    doc = _load(_docx(tmp_path)).documents[0]

    assert "# 年报" in doc.text
    assert "## 财务" in doc.text
    assert doc.metadata["parser"] == "docx"


def test_docx_table_becomes_a_markdown_table(tmp_path: Path) -> None:
    doc = _load(_docx(tmp_path)).documents[0]

    assert "| 项目 | 金额 |" in doc.text
    assert "| 营收 | 100 |" in doc.text


def test_docx_keeps_document_order(tmp_path: Path) -> None:
    """python-docx 把段落和表格分成两个集合，直接遍历会让表格全跑到文末。"""
    text = _load(_docx(tmp_path)).documents[0].text

    assert text.index("| 营收 | 100 |") < text.index("表后的一段。")


def test_localized_word_heading_styles_are_recognized() -> None:
    """中文 Word 存的样式名是「标题 1」而不是「Heading 1」—— 只认英文的话标题层级全丢。"""

    class _Paragraph:
        def __init__(self, name: str) -> None:
            self.style = type("Style", (), {"name": name})()

    assert _heading_level(_Paragraph("Heading 2")) == 2  # type: ignore[arg-type]
    assert _heading_level(_Paragraph("标题 3")) == 3  # type: ignore[arg-type]
    assert _heading_level(_Paragraph("正文")) is None  # type: ignore[arg-type]


def test_a_paragraph_without_a_style_does_not_crash() -> None:
    """`p.style` 可能是 None —— 实测一份 10MB 的文档 890 个段落全是（整篇没用样式）。

    直接取 ``.name`` 会抛 AttributeError，把整份文档变成 skipped，原因还是一句
    ``'NoneType' object has no attribute 'name'`` —— 照着修都没法修。
    """

    class _Paragraph:
        style = None

    assert _heading_level(_Paragraph()) is None  # type: ignore[arg-type]


def _docx_unstyled(tmp_path: Path) -> Path:
    """仿「整篇没用样式」的文档：结构靠字号 + 加粗表达。"""
    target = tmp_path / "unstyled.docx"
    doc = open_docx()

    doc.add_paragraph().add_run("正文一段，十一号字，不加粗。").font.size = Pt(11)

    for text, size in (("第一章、概述", 18), ("1、小节", 14)):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(size)

    # 正文里的加粗强调（字号与正文相同）：不该被当成标题
    emphasis = doc.add_paragraph()
    run = emphasis.add_run("Agent 的核心定义是：它把模型和工具结合起来。")
    run.bold = True
    run.font.size = Pt(11)

    doc.save(target)
    return target


def test_unstyled_docx_infers_headings_from_direct_formatting(tmp_path: Path) -> None:
    """整篇没用样式时，改从字号 + 加粗推断 —— 否则整份文档会变成没有层级的平文本。"""
    text = _load(_docx_unstyled(tmp_path)).documents[0].text

    assert "# 第一章、概述" in text
    assert "## 1、小节" in text, "字号大的级别要更高"


def test_bold_body_text_is_not_mistaken_for_a_heading(tmp_path: Path) -> None:
    """「整段加粗」不足以判定标题 —— 实测一份文档正文里最长的加粗段有 4307 字。

    兜底靠三条同时成立：整段加粗 **且** 比正文大 **且** 够短。
    """
    text = _load(_docx_unstyled(tmp_path)).documents[0].text

    assert "Agent 的核心定义是" in text
    assert "# Agent 的核心定义是" not in text


def test_a_styled_docx_is_not_re_guessed(tmp_path: Path) -> None:
    """作者用了样式就信样式 —— 不去拿字号猜，免得把加粗的正文提拔成标题。"""
    doc = open_docx()
    doc.add_heading("真标题", level=1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("一段又大又粗的正文。")
    run.bold = True
    run.font.size = Pt(20)
    target = tmp_path / "styled.docx"
    doc.save(target)

    text = _load(target).documents[0].text

    assert "# 真标题" in text
    assert "# 一段又大又粗的正文。" not in text


def test_docx_images_become_markers_and_assets(tmp_path: Path) -> None:
    """docx 里的图要和 PDF 那条路一样：本体进资产目录、正文里留原位标记。

    实测一份 10MB 的文档里有 45 张内容图 —— 不取的话它们会**静默消失**。
    """
    png = pymupdf.open()
    png.new_page().insert_text((72, 72), "图", fontsize=20, fontname="china-s")
    image = png[0].get_pixmap().tobytes("png")
    png.close()

    doc = open_docx()
    doc.add_paragraph("图片上面的一段。")
    doc.add_picture(io.BytesIO(image))
    doc.add_paragraph("图片下面的一段。")
    target = tmp_path / "with-image.docx"
    doc.save(target)

    loaded = _load(target).documents[0]

    assert len(loaded.images) == 1
    marker = loaded.images[0].marker
    before, after = loaded.text.split(marker)
    assert "图片上面的一段。" in before
    assert "图片下面的一段。" in after


# ---- html ----­


def _html(tmp_path: Path) -> Path:
    target = tmp_path / "page.html"
    target.write_text(
        """
        <html><head><style>p{color:red}</style><script>var x=1;</script></head>
        <body>
          <nav>站点导航不该进库</nav>
          <h1>产品手册</h1>
          <p>第一段正文。</p>
          <h2>规格</h2>
          <table>
            <tr><th>型号</th><th>容量</th></tr>
            <tr><td>A1</td><td>512</td></tr>
          </table>
          <ul><li>要点一</li><li>要点二</li></ul>
        </body></html>
        """,
        encoding="utf-8",
    )
    return target


def test_html_headings_and_tables_become_markdown(tmp_path: Path) -> None:
    doc = _load(_html(tmp_path)).documents[0]

    assert "# 产品手册" in doc.text
    assert "## 规格" in doc.text
    assert "| 型号 | 容量 |" in doc.text
    assert "| A1 | 512 |" in doc.text
    assert doc.metadata["parser"] == "html"


def test_html_drops_scripts_styles_and_navigation(tmp_path: Path) -> None:
    """脚本样式与站点导航进了库只会变成噪声命中。"""
    text = _load(_html(tmp_path)).documents[0].text

    assert "color:red" not in text
    assert "var x=1" not in text
    assert "站点导航" not in text


def test_html_lists_become_bullets(tmp_path: Path) -> None:
    assert "- 要点一" in _load(_html(tmp_path)).documents[0].text


def _html_edge(tmp_path: Path) -> Path:
    """刻意刁难的 HTML：四种真实世界里常见、但朴素实现会踩的写法。"""
    target = tmp_path / "edge.html"
    target.write_text(
        """
        <html><body>
          <h2>组件</h2>
          <div>只有 div、没有 p 的一段话。</div>
          <pre><code>kubectl get pods -n milvus
kubectl logs -f milvus-0</code></pre>
          <ul>
            <li>etcd
              <ul><li>存元数据</li></ul>
            </li>
            <li>MinIO</li>
          </ul>
          <p>注意：<strong>不要</strong>把 <code>2379</code> 暴露到公网。</p>
          <div><p>div 里套 p，不要取两遍。</p></div>
        </body></html>
        """,
        encoding="utf-8",
    )
    return target


def test_html_div_only_text_is_kept(tmp_path: Path) -> None:
    """现代前端极常见「只有 div 没有 p」—— 只认 p 的话整段会丢。"""
    assert "只有 div、没有 p 的一段话。" in _load(_html_edge(tmp_path)).documents[0].text


def test_html_code_blocks_are_kept_as_fences(tmp_path: Path) -> None:
    """技术文档的代码块：换行与缩进有意义，不能被压平，更不能整块丢掉。"""
    text = _load(_html_edge(tmp_path)).documents[0].text

    assert "```" in text
    assert "kubectl get pods -n milvus\nkubectl logs -f milvus-0" in text


def test_html_nested_lists_are_not_duplicated(tmp_path: Path) -> None:
    """外层 li 的 get_text 会把子项文字吞进去；子项再出一遍就是同一段文字进库三次。"""
    text = _load(_html_edge(tmp_path)).documents[0].text

    assert text.count("存元数据") == 1
    assert text.count("MinIO") == 1


def test_html_inline_tags_do_not_insert_spaces(tmp_path: Path) -> None:
    """`get_text(" ")` 会在行内标签边界插空格，把中文拆坏：「不要把」→「不要 把」。"""
    text = _load(_html_edge(tmp_path)).documents[0].text

    assert "不要把 2379 暴露到公网。" in text


def test_html_container_with_a_paragraph_inside_is_not_taken_twice(tmp_path: Path) -> None:
    """兜底容器只在它内部**没有**块级元素时才算一段。"""
    assert _load(_html_edge(tmp_path)).documents[0].text.count("不要取两遍") == 1


def _html_page(tmp_path: Path) -> Path:
    """仿真一篇带页头噪声的文档页（真实网页的常见形态）。"""
    target = tmp_path / "page.html"
    target.write_text(
        """
        <html><body>
          <div>页头噪声：搜索框、面包屑、筛选器</div>
          <main>
            <h1>Kafka FAQ\u200b</h1>
            <p>正文第一段。</p>
            <ul><li><p>列表项里套了个 p</p></li></ul>
            <h2>标题里带图标字形\uf0c1</h2>
          </main>
          <div>页脚噪声</div>
        </body></html>
        """,
        encoding="utf-8",
    )
    return target


def test_html_prefers_main_and_drops_page_chrome(tmp_path: Path) -> None:
    """真实网页的页头/页脚/侧栏常占去一半文字量，整页扫会把它们灌成检索噪声。"""
    text = _load(_html_page(tmp_path)).documents[0].text

    assert "正文第一段。" in text
    assert "页头噪声" not in text
    assert "页脚噪声" not in text


def test_html_paragraph_inside_a_list_item_is_not_duplicated(tmp_path: Path) -> None:
    """<li> 的文字已经含了它里面的 <p>，再单独取一遍就是同一段进库两次。"""
    assert _load(_html_page(tmp_path)).documents[0].text.count("列表项里套了个 p") == 1


def test_html_icon_font_glyphs_and_zero_width_chars_are_stripped(tmp_path: Path) -> None:
    """文档站的标题尾部常挂着「复制链接」锚点：零宽字符，或图标字体的私用区码点。

    实测 Confluent 的每个标题尾部都是一个 U+F0C1（Font Awesome 的 fa-link）——
    它不是文字，却会跟着进标题路径，让精确匹配失效。
    """
    text = _load(_html_page(tmp_path)).documents[0].text

    assert "# Kafka FAQ" in text
    assert "\u200b" not in text
    assert "\uf0c1" not in text
    assert "## 标题里带图标字形" in text


def test_html_title_decorations_are_stripped(tmp_path: Path) -> None:
    """标题尾部的锚点符与折叠箭头也不算标题的一部分。

    实测那份 Confluent Kafka **Go 客户端**文档（pkg.go.dev 一类的文档站）：313 个标题里
    307 个带着 `¶` 锚点、4 个带 `▹`/`▾` 折叠箭头 —— 不清掉就会整条进 heading_path，
    章节名没法精确匹配，按章节检索也跟着失真。
    """
    page = tmp_path / "go.html"
    page.write_text(
        "<html><body><main>"
        "<h1>Package kafka ¶</h1>"
        "<h2>Overview ▹</h2>"
        "<h3>High-level Consumer ¶</h3>"
        "<p>正文。</p>"
        "</main></body></html>",
        encoding="utf-8",
    )

    text = _load(page).documents[0].text

    assert "# Package kafka" in text
    assert "## Overview" in text
    assert "### High-level Consumer" in text
    for mark in ("¶", "▹", "▾"):
        assert mark not in text


# ---- xlsx ----­


def _xlsx(tmp_path: Path) -> Path:
    target = tmp_path / "book.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "客户清单"
    sheet.append(["名称", "数量"])
    sheet.append(["甲公司", 512])
    sheet.append(["乙公司", 1024.0])
    book.create_sheet("空表")  # 只有表名、没有内容
    book.save(target)
    return target


def test_xlsx_row_becomes_a_self_describing_record(tmp_path: Path) -> None:
    """表名与列名焊进每条记录。

    `| 甲公司 | 512 |` 不知道列名就检索不到 —— 用户问的是「甲公司的数量是多少」，
    答案里必须自带这些词。
    """
    doc = _load(_xlsx(tmp_path)).documents[0]

    assert "客户清单：名称=甲公司，数量=512" in doc.text
    assert doc.metadata["parser"] == "xlsx"


def test_xlsx_integer_like_floats_lose_the_decimal(tmp_path: Path) -> None:
    """Excel 里 512 读出来常是 512.0 —— 留着小数点会让「512」检索不到。"""
    assert "数量=1024" in _load(_xlsx(tmp_path)).documents[0].text


def test_empty_sheet_produces_nothing(tmp_path: Path) -> None:
    assert "空表" not in _load(_xlsx(tmp_path)).documents[0].text


def test_xlsx_declares_row_chunking(tmp_path: Path) -> None:
    """表格是「一行一条记录」，不该按文本结构切 —— 由 loader 声明切法。"""
    assert _load(_xlsx(tmp_path)).documents[0].metadata["chunk_strategy"] == "table-rows"


class _StubEmbedder:
    """散文那条路（markdown-semantic）要用，这里给个最小替身。"""

    @property
    def dim(self) -> int:
        return 2

    async def embed_documents(self, texts: list[str]) -> Embeddings:
        return Embeddings(dense=[[1.0, 0.0] for _ in texts])

    async def embed_query(self, text: str) -> Embeddings:
        return Embeddings(dense=[[1.0, 0.0]])


def _settings() -> Settings:
    return Settings(
        indexing=IndexingConfig(
            chunk=ChunkingConfig(strategy="markdown-semantic", size=1000, overlap=0)
        )
    )


def test_each_row_is_one_chunk_neither_merged_nor_split(tmp_path: Path) -> None:
    """一行就是一条记录：两行并成一块会把记录混在一起，一行切两半则两边都不完整。"""
    prep = prepare(_xlsx(tmp_path), _settings(), None)

    assert [c.text for c in prep.chunks] == [
        "客户清单：名称=甲公司，数量=512",
        "客户清单：名称=乙公司，数量=1024",
    ]


def test_prose_and_tables_take_different_paths_in_one_corpus(tmp_path: Path) -> None:
    """一份语料里散文和表格混着 —— 一个全局策略服务不了两边，所以按文档分发。"""
    _xlsx(tmp_path)
    (tmp_path / "a.md").write_text("# 标题\n\n正文一段。", encoding="utf-8")

    prep = prepare(tmp_path, _settings(), None, _StubEmbedder())  # type: ignore[arg-type]

    table = [c for c in prep.chunks if c.metadata["source_uri"].endswith(".xlsx")]
    prose = [c for c in prep.chunks if c.metadata["source_uri"].endswith(".md")]
    assert len(table) == 2, "表格两行 → 两块"
    assert prose and prose[0].heading_path == "标题", "散文仍走 header 那条路"
    assert prose[0].heading_numbers == "1"


# ---- 注册表 ----­


@pytest.mark.parametrize(
    ("suffix", "name"),
    [(".docx", "docx"), (".html", "html"), (".htm", "html"), (".xlsx", "xlsx"), (".xlsm", "xlsx")],
)
def test_new_suffixes_are_registered(suffix: str, name: str) -> None:
    assert suffix in LOADERS
    assert build_loader(suffix, IngestConfig()).name == name
